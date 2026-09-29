"""Annotate title exposure periods in the established Huawei metrics workbook."""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime
import hashlib
from itertools import chain
import json
from pathlib import Path
import re
import shutil
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel, to_excel


HEADERS = {
    0: "日期",
    1: "图片标题",
    4: "轮播页曝光次数",
    7: "详情页曝光次数",
    13: "标题",
    14: "轮播页曝光次数",
    15: "详情页曝光次数",
}


def parse_day(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(value, str):
        return date.fromisoformat(value.strip().split(" ")[0])
    raise ValueError(f"无法解析明细日期: {value!r}")


def number(value, row_num, column):
    if value is None:
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{column}{row_num} 不是数值: {value!r}")
    return value


def read_records(source: Path, mode: str):
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for index, expected in HEADERS.items():
        if index >= len(header) or expected not in str(header[index] or ""):
            raise ValueError(f"表头不符：{chr(65 + index)} 列应为 {expected}")

    daily = defaultdict(lambda: defaultdict(lambda: [0, 0, False]))
    totals = defaultdict(lambda: [0, 0])
    day_bounds = {}
    titles = []
    seen_details = set()
    duplicate_rows = set()
    previous_day = None
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if len(row) < 16:
            raise ValueError(f"第 {row_num} 行缺少 N/P 列")
        day = parse_day(row[0], book.epoch)
        if previous_day and day < previous_day:
            raise ValueError(f"A 列日期未按升序排列：第 {row_num} 行")
        previous_day = day

        title = row[13]
        if title is not None and str(title).strip():
            if not isinstance(title, str):
                raise ValueError(f"N{row_num} 不是标题文本")
            if any(value is not None for value in row[19:]):
                raise ValueError(f"第 {row_num} 行 S 列之后有内容，不能安全重排")

        detail_key = tuple(row[:12])
        if detail_key in seen_details:
            if title is not None and str(title).strip():
                raise ValueError(f"第 {row_num} 行明细重复且承载 N 列标题，需先人工处理")
            duplicate_rows.add(row_num)
            continue
        seen_details.add(detail_key)
        new_row = row_num - len(duplicate_rows)
        day_bounds.setdefault(day, [new_row, new_row])[1] = new_row
        if title is not None and str(title).strip():
            titles.append((row_num, title, row[16] if len(row) > 16 else None))

        image_title = row[1]
        if image_title is None:
            continue
        e = number(row[4], row_num, "E")
        h = number(row[7], row_num, "H")
        data = daily[image_title][day]
        data[0] += e
        data[1] += h
        data[2] |= h > 0
        totals[image_title][0] += e
        totals[image_title][1] += h

    if not titles or [row for row, _, _ in titles] != list(range(2, len(titles) + 2)):
        raise ValueError("N 列标题必须从第 2 行连续排列")
    if len({title for _, title, _ in titles}) != len(titles):
        raise ValueError("N 列含重复标题，请先明确是否合并")

    records = []
    for old_row, title, q_value in titles:
        details = daily[title]
        positive = sorted(day for day, values in details.items() if values[2])
        start = positive[0] if positive else None
        actual_days = len(positive)
        cutoff = positive[30] if mode == "monthly" and actual_days >= 31 else None
        if cutoff:
            e = sum(v[0] for day, v in details.items() if start <= day <= cutoff)
            h = sum(v[1] for day, v in details.items() if start <= day <= cutoff)
            first_row, last_row = day_bounds[start][0], day_bounds[cutoff][1]
        else:
            e, h = totals[title]
            first_row = last_row = None
        records.append({
            "old_row": old_row,
            "title": title,
            "q_value": q_value,
            "start": start,
            "actual_days": actual_days,
            "shown_days": min(actual_days, 31) if mode == "monthly" else actual_days,
            "cutoff": cutoff,
            "e": e,
            "h": h,
            "first_row": first_row,
            "last_row": last_row,
        })
    book.close()
    if mode == "monthly":
        records.sort(key=lambda record: (-record["h"], record["old_row"]))
    return records, sheet.max_row - len(duplicate_rows), duplicate_rows


def cell(raw: bytes, column: str, row: int) -> bytes:
    match = re.search(
        rb'<c r="' + column.encode() + str(row).encode() + rb'"[^>]*(?:/>|>.*?</c>)',
        raw,
        re.S,
    )
    if not match:
        raise ValueError(f"工作表 XML 缺少 {column}{row}")
    return match.group()


def optional_cell(raw: bytes, column: str, row: int) -> bytes:
    try:
        return cell(raw, column, row)
    except ValueError:
        return b""


def styled_cell(reference: str, style: int, value) -> bytes:
    if isinstance(value, str):
        return (f'<c r="{reference}" s="{style}" t="inlineStr"><is><t>{value}</t></is></c>').encode()
    return f'<c r="{reference}" s="{style}"><v>{value}</v></c>'.encode()


def formula_cell(column: str, row: int, record: dict, source_style: int) -> bytes:
    metric = "E" if column == "O" else "H"
    if record["cutoff"]:
        first, last = record["first_row"], record["last_row"]
        formula = f"SUMIF(B${first}:B${last},N{row},{metric}${first}:{metric}${last})"
    else:
        formula = f"SUMIF(B:B,N{row},{metric}:{metric})"
    value = record["e"] if column == "O" else record["h"]
    return f'<c r="{column}{row}" s="{source_style}"><f>{formula}</f><v>{value}</v></c>'.encode()


def add_styles(xml: bytes):
    used_ids = [int(n) for n in re.findall(rb'numFmtId="(\d+)"', xml)]
    date_id = max([163, *used_ids]) + 1
    days_id = date_id + 1
    custom = (
        f'<numFmt numFmtId="{date_id}" formatCode="m&quot;月&quot;d&quot;日&quot;"/>'
        f'<numFmt numFmtId="{days_id}" formatCode="0&quot;天&quot;"/>'
    ).encode()
    fmt = re.search(rb'<numFmts count="(\d+)">', xml)
    if fmt:
        xml = xml[:fmt.start()] + fmt.group().replace(fmt.group(1), str(int(fmt.group(1)) + 2).encode(), 1) + xml[fmt.end():]
        xml = xml.replace(b'</numFmts>', custom + b'</numFmts>', 1)
    else:
        empty = re.search(rb'<numFmts count="0"\s*/>', xml)
        replacement = b'<numFmts count="2">' + custom + b'</numFmts>'
        if empty:
            xml = xml[:empty.start()] + replacement + xml[empty.end():]
        else:
            xml = xml.replace(b'<fonts ', replacement + b'<fonts ', 1)

    xfs = re.search(rb'<cellXfs count="(\d+)">', xml)
    if not xfs:
        raise ValueError("工作簿缺少 cellXfs 样式表")
    first_style = int(xfs.group(1))
    xml = xml[:xfs.start()] + xfs.group().replace(xfs.group(1), str(first_style + 2).encode(), 1) + xml[xfs.end():]
    base = ('<xf numFmtId="{num}" fontId="0" fillId="0" borderId="0" xfId="0" '
            'applyNumberFormat="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf>')
    xml = xml.replace(b'</cellXfs>',
                      (base.format(num=date_id) + base.format(num=days_id)).encode() + b'</cellXfs>', 1)
    return xml, first_style, first_style + 1


def write_tail(reader, writer, first_chunk: bytes, duplicate_rows: set[int], first_row: int):
    if not duplicate_rows:
        writer.write(first_chunk)
        shutil.copyfileobj(reader, writer, 1 << 20)
        return

    row_pattern = re.compile(rb'<row r="(\d+)"[^>]*>.*?</row>', re.S)
    buffer = b""
    removed = 0
    expected_row = first_row
    for chunk in chain((first_chunk,), iter(lambda: reader.read(1 << 20), b"")):
        buffer += chunk
        position = 0
        for match in row_pattern.finditer(buffer):
            old_row = int(match.group(1))
            if old_row != expected_row:
                raise ValueError(f"明细 XML 行号中断：预期 {expected_row}，实际 {old_row}")
            expected_row += 1
            writer.write(buffer[position:match.start()])
            position = match.end()
            if old_row in duplicate_rows:
                removed += 1
                continue
            new_row = old_row - removed
            raw = match.group()
            if b'<f' in raw and new_row != old_row:
                raise ValueError(f"第 {old_row} 行含公式，无法安全移动明细")
            raw = raw.replace(f'<row r="{old_row}"'.encode(),
                              f'<row r="{new_row}"'.encode(), 1)
            raw = re.sub(
                rb'(<c r="[A-Z]+)' + str(old_row).encode() + rb'"',
                lambda item: item.group(1) + str(new_row).encode() + b'"', raw)
            writer.write(raw)
        buffer = buffer[position:]
    writer.write(buffer)
    if removed != len(duplicate_rows):
        raise ValueError(f"预期删除 {len(duplicate_rows)} 行，实际删除 {removed} 行")


def build_sheet(reader, writer, records, mode, max_row, date_style, days_style, duplicate_rows):
    first_row = 2
    last_title_row = len(records) + 1
    boundary = f'<row r="{last_title_row + 1}"'.encode()
    buffer = bytearray()
    while boundary not in buffer:
        chunk = reader.read(1 << 20)
        if not chunk:
            raise ValueError("找不到标题区后的第一行")
        buffer.extend(chunk)
        if len(buffer) > 20 << 20:
            raise ValueError("标题区 XML 过大，停止定向写回")
    split = buffer.index(boundary)
    early, remainder = bytes(buffer[:split]), bytes(buffer[split:])
    matches = list(re.finditer(rb'<row r="(\d+)"[^>]*>.*?</row>', early, re.S))
    old_rows = {int(m.group(1)): m.group() for m in matches}
    if set(old_rows) != set(range(1, last_title_row + 1)):
        raise ValueError("标题区行号不连续，停止定向写回")

    lead = early[:matches[0].start()]
    dimension = re.search(rb'<dimension ref="([A-Z]+\d+):([A-Z]+\d+)"\s*/>', lead)
    if dimension:
        lead = lead[:dimension.start()] + f'<dimension ref="A1:S{max_row}"/>'.encode() + lead[dimension.end():]
    widths = b'<col min="18" max="18" width="18" customWidth="1"/><col min="19" max="19" width="19" customWidth="1"/>'
    if b'</cols>' in lead:
        lead = lead.replace(b'</cols>', widths + b'</cols>', 1)
    else:
        lead = lead.replace(b'<sheetData>', b'<cols>' + widths + b'</cols><sheetData>', 1)
    writer.write(lead)

    header = old_rows[1]
    style_match = re.search(rb'<c r="N1" s="(\d+)"', header)
    header_style = int(style_match.group(1)) if style_match else 0
    header = re.sub(rb'spans="\d+:\d+"', b'spans="1:19"', header, count=1)
    header = re.sub(rb'<c r="[RS]1"[^>]*(?:/>|>.*?</c>)', b'', header)
    day_header = "本次统计曝光天数" if mode == "monthly" else "累计正曝光天数"
    addition = styled_cell("R1", header_style, "首次正曝光日期") + styled_cell("S1", header_style, day_header)
    writer.write(header.replace(b'</row>', addition + b'</row>', 1))

    for new_row, record in enumerate(records, first_row):
        original_row = old_rows[new_row]
        n_start = re.search(rb'<c r="N' + str(new_row).encode() + rb'"', original_row)
        if not n_start:
            raise ValueError(f"XML 中缺少 N{new_row}")
        left = re.sub(rb'spans="\d+:\d+"', b'spans="1:19"', original_row[:n_start.start()], count=1)
        if mode == "monthly":
            old_row = record["old_row"]
            n = cell(old_rows[old_row], "N", old_row).replace(
                f'r="N{old_row}"'.encode(), f'r="N{new_row}"'.encode(), 1)
            q = optional_cell(old_rows[old_row], "Q", old_row).replace(
                f'r="Q{old_row}"'.encode(), f'r="Q{new_row}"'.encode(), 1)
            o_xml = cell(original_row, "O", new_row)
            p_xml = cell(original_row, "P", new_row)
            o_style = re.search(rb's="(\d+)"', o_xml)
            p_style = re.search(rb's="(\d+)"', p_xml)
            summary = (n + formula_cell("O", new_row, record, int(o_style.group(1)) if o_style else 0)
                       + formula_cell("P", new_row, record, int(p_style.group(1)) if p_style else 0)
                       + q)
        else:
            summary = cell(original_row, "N", new_row)
            if duplicate_rows:
                for column in ("O", "P"):
                    original_cell = cell(original_row, column, new_row)
                    style = re.search(rb's="(\d+)"', original_cell)
                    summary += formula_cell(column, new_row, record,
                                            int(style.group(1)) if style else 0)
            else:
                summary += cell(original_row, "O", new_row) + cell(original_row, "P", new_row)
            summary += optional_cell(original_row, "Q", new_row)
        start = record["start"]
        r_value = int(to_excel(start)) if start else "无正曝光数据"
        r_style = date_style if start else 0
        row_xml = (left + summary + styled_cell(f"R{new_row}", r_style, r_value)
                   + styled_cell(f"S{new_row}", days_style, record["shown_days"]) + b'</row>')
        writer.write(row_xml)

    write_tail(reader, writer, remainder, duplicate_rows, last_title_row + 1)


def write_workbook(source: Path, target: Path, records, mode: str, max_row: int,
                   duplicate_rows: set[int]):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, "w") as output:
        style_xml, date_style, days_style = add_styles(original.read("xl/styles.xml"))
        for item in original.infolist():
            if item.filename == "xl/styles.xml":
                output.writestr(item, style_xml)
            elif item.filename == "xl/worksheets/sheet1.xml":
                with original.open(item) as reader, output.open(item, "w") as writer:
                    build_sheet(reader, writer, records, mode, max_row, date_style,
                                days_style, duplicate_rows)
            else:
                output.writestr(item, original.read(item))


def verify(source: Path, target: Path, records, mode: str, max_row: int,
           duplicate_rows: set[int]):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target) as output:
        if output.testzip() is not None:
            raise ValueError("输出 XLSX ZIP 校验失败")
        if set(original.namelist()) != set(output.namelist()):
            raise ValueError("输出工作簿部件数量改变")
        changed = {name for name in original.namelist()
                   if hashlib.sha256(original.read(name)).digest() != hashlib.sha256(output.read(name)).digest()}
        if changed != {"xl/styles.xml", "xl/worksheets/sheet1.xml"}:
            raise ValueError(f"改动了非目标工作簿部件: {changed}")
    book = load_workbook(target, read_only=True, data_only=True)
    if book.worksheets[0].max_row != max_row:
        raise ValueError("输出明细行数与去重预期不符")
    rows = book.worksheets[0].iter_rows(min_row=2, max_row=len(records) + 1, values_only=True)
    for record, row in zip(records, rows):
        if row[13] != record["title"] or row[18] != record["shown_days"]:
            raise ValueError("输出标题或 S 列与独立统计不符")
        if row[16] != record["q_value"]:
            raise ValueError("Q 列标记未随标题正确移动")
        if record["start"]:
            r = row[17]
            if not isinstance(r, (datetime, date)) or (r.date() if isinstance(r, datetime) else r) != record["start"]:
                raise ValueError("输出 R 列日期不符")
        elif row[17] != "无正曝光数据":
            raise ValueError("无正曝光标题的 R 列标记不符")
        if (mode == "monthly" or duplicate_rows) and (row[14], row[15]) != (record["e"], record["h"]):
            raise ValueError("输出 O/P 汇总值不符")
    if duplicate_rows:
        seen = set()
        for row in book.worksheets[0].iter_rows(min_row=2, max_col=12, values_only=True):
            key = tuple(row)
            if key in seen:
                raise ValueError("输出 A–L 仍有完全重复的明细")
            seen.add(key)
    book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--mode", choices=("fortnight", "monthly"), required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    source, target = args.input.resolve(), args.output.resolve()
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        parser.error("输入文件必须是存在的 .xlsx")
    if source == target:
        parser.error("必须另存副本，不能覆盖源文件")
    if target.exists():
        parser.error("输出文件已存在，请使用新的副本路径")

    records, max_row, duplicate_rows = read_records(source, args.mode)
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        write_workbook(source, target, records, args.mode, max_row, duplicate_rows)
        verify(source, target, records, args.mode, max_row, duplicate_rows)
    except Exception:
        target.unlink(missing_ok=True)
        raise
    report = {
        "output": str(target), "mode": args.mode, "titles": len(records),
        "removed_duplicate_rows": len(duplicate_rows),
        "positive": sum(r["actual_days"] > 0 for r in records),
        "zero": sum(r["actual_days"] == 0 for r in records),
    }
    if args.mode == "monthly":
        report.update(capped=sum(r["actual_days"] >= 31 for r in records),
                      under31=sum(r["actual_days"] < 31 for r in records))
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
